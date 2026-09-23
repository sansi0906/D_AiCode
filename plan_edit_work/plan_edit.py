# -*- coding: utf-8 -*-
"""精简《生活服务小程序上架方案》：只保留上架准备工作相关内容"""
import copy, re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PATH = r'E:\DAI\生活服务小程序上架方案.docx'
doc = Document(PATH)
body = doc.element.body

paras = list(doc.paragraphs)  # 快照


def ptext(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def is_h(style_name):
    """返回样式名为指定值的正文标题元素列表（排除目录条目：含制表符页码）"""
    out = []
    for p in paras:
        if p.style.name == style_name and '\t' not in p.text:
            out.append(p._element)
    return out


def set_para_text(p_el, text):
    runs = p_el.findall(qn('w:r'))
    for i, r in enumerate(runs):
        for t in r.findall(qn('w:t')):
            r.remove(t)
        if i == 0:
            t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
            r.append(t)


# ---------- 1. 提取 8.2 内容（删除八章前） ----------
h2s = is_h('Heading 2')
extract = []
for i, el in enumerate(h2s):
    if ptext(el).strip().startswith('8.2 '):
        extract.append(el)
        # 后续正文段落（非标题）直到下一个 H2/H1
        nxt = el.getnext()
        while nxt is not None:
            if nxt.tag == qn('w:p') and nxt.find(qn('w:pPr')) is not None:
                ps = nxt.find(qn('w:pPr')).find(qn('w:pStyle'))
                if ps is not None:
                    sid = ps.get(qn('w:val'))
                    if sid in ('Heading2', 'Heading 2', '4') or sid in ('Heading1', 'Heading 1', '3'):
                        break
            if nxt.tag == qn('w:tbl'):
                break
            extract.append(nxt)
            nxt = nxt.getnext()
        break
for el in extract:
    body.remove(el)

# ---------- 2. 删除无关章节 ----------
h1s = is_h('Heading 1')
def find_h1(prefix):
    for el in h1s:
        if ptext(el).strip().startswith(prefix):
            return el
    return None

def delete_range(start_el, end_el):
    el = start_el
    while el is not None and el is not end_el:
        nxt = el.getnext()
        body.remove(el)
        el = nxt

sec_two = find_h1('二、商业模式设计')
sec_three = find_h1('三、主体与资质准备')
sec_six = find_h1('六、产品与功能规划')
sec_seven = find_h1('七、合规要点')
sec_eight = find_h1('八、开发与测试计划')
sec_nine = find_h1('九、提审与发布')
sec_ten = find_h1('十、上线后运营')
sec_eleven = find_h1('十一、总体时间表')

delete_range(sec_two, sec_three)
delete_range(sec_six, sec_seven)
delete_range(sec_eight, sec_nine)
delete_range(sec_ten, sec_eleven)

# ---------- 3. 删除时间表中开发/提审/上线相关行 ----------
for tbl in doc.tables:
    if '第1周' in ptext(tbl._tbl) and '灰度发布' in ptext(tbl._tbl):
        for tr in list(tbl._tbl.findall(qn('w:tr'))):
            rowtxt = ptext(tr)
            if any(k in rowtxt for k in ['第3～8周', '第8周', '第9周', '第10周']):
                tbl._tbl.remove(tr)
        break

# ---------- 4. 删除 13.5 中开发/上线相关勾选项 ----------
targets = ['□ 体验版全流程走查通过', '□ 代码审核通过',
           '□ 灰度发布完成并观察指标', '□ 正式全量上线']
for p in paras:
    if p.text.strip() in targets:
        p._element.getparent().remove(p._element)

# ---------- 5. 标题重编号 ----------
h1_map = {
    '三、主体与资质准备': '二、主体与资质准备',
    '四、备案与网络合规': '三、备案与网络合规',
    '五、支付方案': '四、支付开通准备',
    '七、合规要点': '五、合规要点',
    '九、提审与发布': '七、提审材料准备',
    '十一、总体时间表（建议排期，以实际启动日顺延）': '八、准备阶段时间表（建议排期，以实际启动日顺延）',
    '十二、风险与应对': '九、风险与应对',
    '十三、附录：全流程检查清单': '十、附录：上架准备检查清单',
}
h2_renum = {
    '3.1': '2.1', '3.2': '2.2', '3.3': '2.3', '3.4': '2.4', '3.5': '2.5',
    '4.1': '3.1', '4.2': '3.2', '4.3': '3.3',
    '5.1': '4.1', '5.2': '4.2', '5.3': '4.3',
    '7.1': '5.1', '7.2': '5.2', '7.3': '5.3', '7.4': '5.4', '7.5': '5.5',
    '9.1': '7.1', '9.2': '7.2',
    '13.1': '10.1', '13.2': '10.2', '13.3': '10.3', '13.4': '10.4',
    '13.5': '10.5',
}
for el in h1s:
    txt = ptext(el).strip()
    if txt in h1_map:
        set_para_text(el, h1_map[txt])
for el in h2s:
    txt = ptext(el).strip()
    m = re.match(r'^(\d+\.\d+)\s', txt)
    if m and m.group(1) in h2_renum:
        set_para_text(el, h2_renum[m.group(1)] + txt[m.end():])

# ---------- 6. 插入新章节六（账号与开发环境准备） ----------
target = None
for el in h1s:
    if ptext(el).strip().startswith('七、提审材料准备'):
        target = el
        break

tpl = h1s[0] if h1s else None
new_h1 = copy.deepcopy(tpl)
set_para_text(new_h1, '六、账号与开发环境准备')
target.addprevious(new_h1)

set_para_text(extract[0], '6.1 环境与账号准备清单')
for el in reversed(extract):
    target.addprevious(el)

# ---------- 7. 交叉引用替换 ----------
repl = [
    ('见 3.1 建议文本', '见 2.1 建议文本'),
    ('见 3.4', '见 2.4'),
    ('（见3.4）', '（见2.4）'),
    ('见第五章', '见第四章'),
    ('按3.3矩阵逐类核对', '按2.3矩阵逐类核对'),
]
for t in body.iter(qn('w:t')):
    if t.text:
        for a, b in repl:
            if a in t.text:
                t.text = t.text.replace(a, b)

doc.save(PATH)
print('EDIT_OK')
